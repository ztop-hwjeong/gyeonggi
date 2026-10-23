from copy import copy
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def test_render_expands_rows_translates_formulas_and_preserves_layout(
    client: TestClient, report_input: dict, template_directory: Path
):
    template = template_directory / "education-audit-press" / "1.0.0" / "template.xlsx"
    original = template.read_bytes()
    report_input["issues"] *= 30
    report_input["articles"] *= 75
    response = client.post(
        "/templates/education-audit-press/versions/1.0.0/render", json={"data": report_input}
    )
    assert response.status_code == 200, response.text
    assert template.read_bytes() == original
    assert response.headers["content-type"].endswith("spreadsheetml.sheet")
    source = load_workbook(BytesIO(original))
    result = load_workbook(BytesIO(response.content))
    summary, articles = result.worksheets
    assert summary["B43"].value == "정책 이슈"
    assert "C43" in summary["D43"].value
    assert "C14" not in summary["D43"].value
    assert articles["D76"].value == "기사"
    assert articles["A2"].value == 1
    assert "2026-10-29" in summary["B3"].value
    assert "{{" not in summary["B3"].value
    for sheet, prototype, addresses in [
        (summary, source.worksheets[0], ["B14", "C14", "D14", "B43", "C43", "D43"]),
        (articles, source.worksheets[1], ["A2", "D2", "F2", "A76", "D76", "F76"]),
    ]:
        for address in addresses:
            prototype_row = 14 if sheet is summary else 2
            expected = prototype[f"{sheet[address].column_letter}{prototype_row}"]
            for attribute in ("font", "fill", "border", "alignment", "number_format", "protection"):
                assert copy(getattr(sheet[address], attribute)) == copy(
                    getattr(expected, attribute)
                )
        for column in prototype.column_dimensions:
            assert (
                sheet.column_dimensions[column].width == prototype.column_dimensions[column].width
            )
        assert sheet.page_margins == prototype.page_margins
    assert articles.row_dimensions[76].height == source.worksheets[1].row_dimensions[2].height
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(response.content)) as after:
        for part in before.namelist():
            if part.startswith("xl/printerSettings/"):
                assert after.read(part) == before.read(part)
        for part in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
            old = ElementTree.fromstring(before.read(part)).find("{*}pageSetup")
            new = ElementTree.fromstring(after.read(part)).find("{*}pageSetup")
            assert old is not None and new is not None
            assert old.attrib == new.attrib


def test_render_accepts_empty_tables(client: TestClient):
    response = client.post(
        "/templates/education-audit-press/versions/1.0.0/render",
        json={"data": {"analysis_date": "2026-10-29", "issues": [], "articles": []}},
    )
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.worksheets[1]["A2"].value is None


def test_render_rejects_invalid_input_before_generation(client: TestClient, report_input: dict):
    report_input["analysis_date"] = "2026-02-30"
    response = client.post(
        "/templates/education-audit-press/versions/1.0.0/render", json={"data": report_input}
    )
    assert response.status_code == 422
    assert response.json()["detail"][0]["loc"] == ["body", "data", "analysis_date"]


def test_render_missing_version_returns_404(client: TestClient, report_input: dict):
    response = client.post(
        "/templates/education-audit-press/versions/9.0.0/render", json={"data": report_input}
    )
    assert response.status_code == 404


def test_article_text_is_not_executed_as_formula(client: TestClient, report_input: dict):
    report_input["articles"][0]["title"] = "=1+1"
    response = client.post(
        "/templates/education-audit-press/versions/1.0.0/render", json={"data": report_input}
    )
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    cell = workbook.worksheets[1]["D2"]
    assert cell.value == "=1+1"
    assert cell.data_type == "s"
